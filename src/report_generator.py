"""
Genera reportes CSV y Excel con anomalias detectadas.
Incluye imagenes recortadas de los slots anomalos en el XLSX.
"""

import os
import io
import csv
import json
from datetime import datetime
from pathlib import Path
from dataclasses import dataclass, field

import cv2
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .anomaly_detector import SlotResult
from .roi_extractor import FieldROI


CLASSES = ["0", "1", "2", "3", "4", "5", "6", "7", "8", "9", "dot", "lines"]


@dataclass
class AnomalyRecord:
    pdf_name:        str
    page:            int
    section:         str
    list_name:       str
    field_name:      str
    candidate:       int | None
    col_id:          str | None
    digit_slot:      int
    predicted:       str
    confidence:      float
    entropy:         float
    margin:          float
    top3:            list[tuple[str, float]]
    status:          str
    reasons:         list[str]
    crop_path:       str
    crop_img:        np.ndarray = field(default=None, repr=False)


def build_records(pdf_name: str,
                  rois: list[FieldROI],
                  results: list[SlotResult],
                  crops_dir: str) -> list[AnomalyRecord]:
    records = []
    for roi, res in zip(rois, results):
        if not res.is_anomaly and res.is_blank:
            continue

        crop_filename = (
            f"{Path(pdf_name).stem}_p{roi.page:02d}_{roi.field_name}"
            f"_s{roi.digit_slot}.png"
        )
        crop_path = os.path.join(crops_dir, crop_filename)

        if roi.crop is not None and roi.crop.size > 0:
            cv2.imwrite(crop_path, roi.crop)

        records.append(AnomalyRecord(
            pdf_name   = pdf_name,
            page       = roi.page,
            section    = roi.section,
            list_name  = roi.list_name,
            field_name = roi.field_name,
            candidate  = roi.candidate,
            col_id     = roi.col_id,
            digit_slot = roi.digit_slot,
            predicted  = res.predicted_label,
            confidence = res.confidence,
            entropy    = res.entropy,
            margin     = res.margin,
            top3       = res.top3,
            status     = "ANOMALIA" if res.is_anomaly else "OK",
            reasons    = res.anomaly_reasons,
            crop_path  = crop_path,
            crop_img   = roi.crop,
        ))

    return records


def _records_to_dicts(records: list[AnomalyRecord]) -> list[dict]:
    rows = []
    for r in records:
        top3_str = " | ".join(f"{cls}:{p:.2%}" for cls, p in r.top3)
        rows.append({
            "archivo":     r.pdf_name,
            "pagina":      r.page,
            "seccion":     r.section,
            "lista":       r.list_name,
            "campo":       r.field_name,
            "candidato":   r.candidate if r.candidate is not None else "",
            "columna":     r.col_id if r.col_id else "",
            "slot_digito": r.digit_slot,
            "prediccion":  r.predicted,
            "confianza":   f"{r.confidence:.2%}",
            "entropia":    f"{r.entropy:.4f}",
            "margen":      f"{r.margin:.4f}",
            "top3":        top3_str,
            "estado":      r.status,
            "motivo":      "; ".join(r.reasons),
            "crop":        r.crop_path,
        })
    return rows


def export_csv(records: list[AnomalyRecord], out_path: str):
    rows = _records_to_dicts(records)
    if not rows:
        print("[INFO] Sin anomalias para exportar al CSV.")
        return

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"[CSV] {len(rows)} registros -> {out_path}")


def export_xlsx(records: list[AnomalyRecord], out_path: str):
    rows = _records_to_dicts(records)
    if not rows:
        print("[INFO] Sin anomalias para exportar al XLSX.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalias"

    RED_FILL  = PatternFill("solid", fgColor="FFCCCC")
    YEL_FILL  = PatternFill("solid", fgColor="FFFACC")
    HDR_FILL  = PatternFill("solid", fgColor="1F3864")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
    BOLD_FONT = Font(bold=True, size=9)
    NORM_FONT = Font(size=9)
    CENTER    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT      = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    THIN      = Side(style="thin", color="AAAAAA")
    BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    cols_no_img = list(rows[0].keys())
    cols_no_img.remove("crop")
    headers = cols_no_img + ["imagen"]

    for col_i, hdr in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_i, value=hdr.upper())
        cell.fill    = HDR_FILL
        cell.font    = HDR_FONT
        cell.alignment = CENTER
        cell.border  = BORDER

    ws.row_dimensions[1].height = 22

    col_widths = {
        "archivo": 28, "pagina": 7, "seccion": 22, "lista": 24,
        "campo": 20, "candidato": 9, "columna": 8, "slot_digito": 8,
        "prediccion": 9, "confianza": 9, "entropia": 9, "margen": 9,
        "top3": 30, "estado": 10, "motivo": 32, "imagen": 14,
    }
    for col_i, hdr in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_i)].width = col_widths.get(hdr, 12)

    IMG_ROW_H = 55
    IMG_COL   = len(headers)

    for row_i, (rec, row_dict) in enumerate(zip(records, rows), start=2):
        fill = RED_FILL if rec.status == "ANOMALIA" else YEL_FILL

        for col_i, hdr in enumerate(cols_no_img, start=1):
            cell = ws.cell(row=row_i, column=col_i, value=row_dict[hdr])
            cell.fill      = fill
            cell.font      = BOLD_FONT if hdr == "estado" else NORM_FONT
            cell.alignment = CENTER if hdr in ("pagina", "slot_digito", "prediccion",
                                                "confianza", "estado", "candidato") else LEFT
            cell.border = BORDER

        ws.row_dimensions[row_i].height = IMG_ROW_H

        if rec.crop_img is not None and rec.crop_img.size > 0:
            try:
                big = cv2.resize(rec.crop_img, (80, 60), interpolation=cv2.INTER_NEAREST)
                _, buf = cv2.imencode(".png", big)
                xl_img = XLImage(io.BytesIO(buf.tobytes()))
                xl_img.width  = 72
                xl_img.height = 54
                cell_addr = f"{get_column_letter(IMG_COL)}{row_i}"
                ws.add_image(xl_img, cell_addr)
            except Exception:
                pass

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    summary_ws = wb.create_sheet("Resumen")
    from collections import Counter
    by_pdf  = Counter(r.pdf_name for r in records if r.status == "ANOMALIA")
    by_page = Counter((r.pdf_name, r.page) for r in records if r.status == "ANOMALIA")

    summary_ws.append(["RESUMEN DE ANOMALIAS"])
    summary_ws.append([])
    summary_ws.append(["Total anomalias", sum(1 for r in records if r.status == "ANOMALIA")])
    summary_ws.append(["Total registros", len(records)])
    summary_ws.append([])
    summary_ws.append(["Archivo", "Anomalias"])
    for pdf, cnt in sorted(by_pdf.items()):
        summary_ws.append([pdf, cnt])

    wb.save(out_path)
    print(f"[XLSX] {len(rows)} registros -> {out_path}")


def export_all(records: list[AnomalyRecord],
               output_dir: str,
               prefix: str = "reporte"):
    ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path  = os.path.join(output_dir, f"{prefix}_{ts}.csv")
    xlsx_path = os.path.join(output_dir, f"{prefix}_{ts}.xlsx")
    export_csv(records,  csv_path)
    export_xlsx(records, xlsx_path)
    return csv_path, xlsx_path
